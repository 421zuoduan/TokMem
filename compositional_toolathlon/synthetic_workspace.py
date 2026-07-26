from __future__ import annotations

import base64
import csv
import hashlib
import json
import re
import shutil
import tempfile
import zipfile
from copy import copy
from datetime import datetime
from pathlib import Path, PurePosixPath
from typing import Any


SUPPORTED_FILE_FORMATS = {"text", "json", "csv", "xlsx", "pdf", "base64"}
SUPPORTED_ASSERTIONS = {
    "exists",
    "absent",
    "text_equals",
    "text_contains",
    "json_equals",
    "csv_equals",
    "xlsx_cells_equal",
    "xlsx_nonempty_cells_equal",
    "xlsx_sheet_names_equal",
    "xlsx_table_names_equal",
    "xlsx_table_properties",
    "xlsx_table_definitions_equal",
    "xlsx_chart_count",
    "xlsx_chart_properties",
    "xlsx_cell_style",
    "xlsx_merged_ranges_equal",
    "pdf_text_contains",
    "pdf_page_count",
    "pdf_page_text_contains",
    "pdf_page_text_equals",
    "sha256",
    "file_size_at_least",
    "tree_equals",
}
DETERMINISTIC_OFFICE_TIMESTAMP = datetime(2000, 1, 1, 0, 0, 0)
DETERMINISTIC_ZIP_TIMESTAMP = (2000, 1, 1, 0, 0, 0)
CELL_COORDINATE_PATTERN = re.compile(r"^[A-Za-z]{1,3}[1-9][0-9]*$")
HEX_COLOR_PATTERN = re.compile(r"^(?:#?[0-9A-Fa-f]{6}|#?[0-9A-Fa-f]{8})$")
XLSX_STYLE_PROPERTIES = {"bold", "font_color", "bg_color", "alignment"}
XLSX_CHART_TYPES = {"area", "bar", "line", "pie"}
XLSX_CHART_ASSERTION_PROPERTIES = {
    "chart_type",
    "target_cell",
    "title",
    "x_axis",
    "y_axis",
    "data_range",
}


def safe_relative_path(value: Any) -> Path:
    if not isinstance(value, str) or not value:
        raise ValueError("workspace path must be a non-empty string")
    if "\\" in value or "\x00" in value:
        raise ValueError(f"workspace path contains a forbidden character: {value!r}")
    path = PurePosixPath(value)
    if path.is_absolute() or any(part in {"", ".", ".."} for part in path.parts):
        raise ValueError(f"workspace path must be normalized and relative: {value!r}")
    return Path(*path.parts)


def _cell_coordinate(value: Any, label: str) -> tuple[int, int]:
    if not isinstance(value, str) or CELL_COORDINATE_PATTERN.fullmatch(value) is None:
        raise ValueError(f"{label} must be a local A1-style cell coordinate")
    match = re.fullmatch(r"([A-Za-z]+)([0-9]+)", value)
    if match is None:
        raise AssertionError("validated cell coordinate did not parse")
    column = 0
    for character in match.group(1).upper():
        column = column * 26 + ord(character) - ord("A") + 1
    row = int(match.group(2))
    if column > 16_384 or row > 1_048_576:
        raise ValueError(f"{label} lies outside Excel worksheet bounds")
    return column, row


def _cell_range(value: Any, label: str) -> tuple[int, int, int, int]:
    if not isinstance(value, str) or value.count(":") != 1:
        raise ValueError(f"{label} must be a local A1-style cell range")
    start, end = value.split(":")
    min_column, min_row = _cell_coordinate(start, f"{label} start")
    max_column, max_row = _cell_coordinate(end, f"{label} end")
    if min_column > max_column or min_row > max_row:
        raise ValueError(f"{label} must run from its top-left to bottom-right cell")
    return min_column, min_row, max_column, max_row


def _column_name(column: int) -> str:
    characters = []
    while column:
        column, remainder = divmod(column - 1, 26)
        characters.append(chr(ord("A") + remainder))
    return "".join(reversed(characters))


def _canonical_cell_coordinate(value: Any, label: str) -> str:
    column, row = _cell_coordinate(value, label)
    return f"{_column_name(column)}{row}"


def _canonical_cell_range(value: Any, label: str) -> str:
    min_column, min_row, max_column, max_row = _cell_range(value, label)
    return (
        f"{_column_name(min_column)}{min_row}:"
        f"{_column_name(max_column)}{max_row}"
    )


def _normalize_hex_color(value: Any, label: str) -> str:
    if not isinstance(value, str) or HEX_COLOR_PATTERN.fullmatch(value) is None:
        raise ValueError(f"{label} must be a 6- or 8-digit hexadecimal color")
    normalized = value.removeprefix("#").upper()
    return normalized[-6:]


def _validate_style_properties(properties: Any, label: str) -> dict[str, Any]:
    if not isinstance(properties, dict) or not properties:
        raise ValueError(f"{label} must be a non-empty object")
    unexpected = sorted(set(properties) - XLSX_STYLE_PROPERTIES)
    if unexpected:
        raise ValueError(f"{label} has unsupported properties: {unexpected}")
    if "bold" in properties and not isinstance(properties["bold"], bool):
        raise ValueError(f"{label}.bold must be boolean")
    for color_property in ("font_color", "bg_color"):
        if color_property in properties:
            _normalize_hex_color(
                properties[color_property],
                f"{label}.{color_property}",
            )
    if "alignment" in properties and (
        not isinstance(properties["alignment"], str)
        or not properties["alignment"].strip()
    ):
        raise ValueError(f"{label}.alignment must be a non-empty string")
    return properties


def _validate_xlsx_content(content: Any) -> dict[str, Any]:
    if not isinstance(content, dict) or set(content) != {"sheets"}:
        raise ValueError("xlsx content requires exactly a sheets list")
    sheets = content["sheets"]
    if not isinstance(sheets, list) or not sheets:
        raise ValueError("xlsx content requires a non-empty sheets list")
    seen_sheet_names: set[str] = set()
    seen_table_names: set[str] = set()
    for sheet_index, sheet_spec in enumerate(sheets):
        label = f"xlsx sheet {sheet_index}"
        if not isinstance(sheet_spec, dict):
            raise ValueError(f"{label} must be an object")
        unexpected = sorted(
            set(sheet_spec)
            - {"name", "rows", "styles", "tables", "charts", "merged_ranges"}
        )
        if unexpected:
            raise ValueError(f"{label} has unsupported keys: {unexpected}")
        title = sheet_spec.get("name")
        rows = sheet_spec.get("rows")
        if not isinstance(title, str) or not title:
            raise ValueError(f"{label} requires a non-empty string name")
        if title in seen_sheet_names:
            raise ValueError(f"xlsx content duplicates sheet name {title!r}")
        seen_sheet_names.add(title)
        if not isinstance(rows, list) or any(not isinstance(row, list) for row in rows):
            raise ValueError(f"{label}.rows must be a list of row lists")

        merged_ranges = sheet_spec.get("merged_ranges", [])
        if not isinstance(merged_ranges, list):
            raise ValueError(f"{label}.merged_ranges must be a list")
        merged_rectangles = [
            _cell_range(merged_range, f"{label}.merged_ranges")
            for merged_range in merged_ranges
        ]
        canonical_merged_ranges = [
            _canonical_cell_range(merged_range, f"{label}.merged_ranges")
            for merged_range in merged_ranges
        ]
        if len(canonical_merged_ranges) != len(set(canonical_merged_ranges)):
            raise ValueError(f"{label}.merged_ranges must contain unique ranges")
        for index, rectangle in enumerate(merged_rectangles):
            min_column, min_row, max_column, max_row = rectangle
            for earlier in merged_rectangles[:index]:
                (
                    earlier_min_column,
                    earlier_min_row,
                    earlier_max_column,
                    earlier_max_row,
                ) = earlier
                overlaps = (
                    max(min_column, earlier_min_column)
                    <= min(max_column, earlier_max_column)
                    and max(min_row, earlier_min_row)
                    <= min(max_row, earlier_max_row)
                )
                if overlaps:
                    raise ValueError(f"{label}.merged_ranges must not overlap")

        styles = sheet_spec.get("styles", [])
        if not isinstance(styles, list):
            raise ValueError(f"{label}.styles must be a list")
        for style_index, style_spec in enumerate(styles):
            style_label = f"{label}.styles[{style_index}]"
            if not isinstance(style_spec, dict):
                raise ValueError(f"{style_label} must be an object")
            unexpected = sorted(
                set(style_spec)
                - {"start_cell", "end_cell", *XLSX_STYLE_PROPERTIES}
            )
            if unexpected:
                raise ValueError(f"{style_label} has unsupported keys: {unexpected}")
            start_column, start_row = _cell_coordinate(
                style_spec.get("start_cell"),
                f"{style_label}.start_cell",
            )
            end_column, end_row = _cell_coordinate(
                style_spec.get("end_cell"),
                f"{style_label}.end_cell",
            )
            if start_column > end_column or start_row > end_row:
                raise ValueError(
                    f"{style_label} must run from its top-left to bottom-right cell"
                )
            _validate_style_properties(
                {
                    key: value
                    for key, value in style_spec.items()
                    if key in XLSX_STYLE_PROPERTIES
                },
                style_label,
            )

        tables = sheet_spec.get("tables", [])
        if not isinstance(tables, list):
            raise ValueError(f"{label}.tables must be a list")
        for table_index, table_spec in enumerate(tables):
            table_label = f"{label}.tables[{table_index}]"
            if not isinstance(table_spec, dict) or set(table_spec) != {
                "data_range",
                "table_name",
                "table_style",
            }:
                raise ValueError(
                    f"{table_label} requires exactly data_range, table_name, "
                    "and table_style"
                )
            _cell_range(table_spec["data_range"], f"{table_label}.data_range")
            table_name = table_spec["table_name"]
            table_style = table_spec["table_style"]
            if (
                not isinstance(table_name, str)
                or not table_name
                or any(character.isspace() for character in table_name)
            ):
                raise ValueError(
                    f"{table_label}.table_name must be non-empty and contain no whitespace"
                )
            if table_name in seen_table_names:
                raise ValueError(f"xlsx content duplicates table name {table_name!r}")
            seen_table_names.add(table_name)
            if not isinstance(table_style, str) or not table_style:
                raise ValueError(f"{table_label}.table_style must be non-empty")

        charts = sheet_spec.get("charts", [])
        if not isinstance(charts, list):
            raise ValueError(f"{label}.charts must be a list")
        for chart_index, chart_spec in enumerate(charts):
            chart_label = f"{label}.charts[{chart_index}]"
            if not isinstance(chart_spec, dict):
                raise ValueError(f"{chart_label} must be an object")
            required = {"data_range", "chart_type", "target_cell"}
            allowed = required | {"title", "x_axis", "y_axis"}
            missing = sorted(required - set(chart_spec))
            unexpected = sorted(set(chart_spec) - allowed)
            if missing or unexpected:
                raise ValueError(
                    f"{chart_label} missing={missing} unsupported={unexpected}"
                )
            _cell_range(chart_spec["data_range"], f"{chart_label}.data_range")
            _cell_coordinate(chart_spec["target_cell"], f"{chart_label}.target_cell")
            if chart_spec["chart_type"] not in XLSX_CHART_TYPES:
                raise ValueError(
                    f"{chart_label}.chart_type must be one of "
                    f"{sorted(XLSX_CHART_TYPES)}"
                )
            for title_key in ("title", "x_axis", "y_axis"):
                if title_key in chart_spec and not isinstance(
                    chart_spec[title_key],
                    str,
                ):
                    raise ValueError(f"{chart_label}.{title_key} must be a string")
    return content


def _validate_assertion_arguments(assertion: dict[str, Any]) -> None:
    operation = assertion["op"]
    if operation in {"xlsx_cells_equal", "xlsx_nonempty_cells_equal"}:
        if set(assertion) != {"op", "path", "sheet", "cells"}:
            raise ValueError(
                f"{operation} requires exactly op, path, sheet, and cells"
            )
        if not isinstance(assertion["sheet"], str) or not assertion["sheet"]:
            raise ValueError(f"{operation} requires a non-empty sheet")
        cells = assertion["cells"]
        if not isinstance(cells, dict):
            raise ValueError(f"{operation}.cells must be an object")
        canonical_cells = [
            _canonical_cell_coordinate(coordinate, f"{operation}.cells")
            for coordinate in cells
        ]
        if len(canonical_cells) != len(set(canonical_cells)):
            raise ValueError(
                f"{operation}.cells must not repeat coordinates by case"
            )
    elif operation == "xlsx_sheet_names_equal":
        sheets = assertion.get("sheets")
        if (
            not isinstance(sheets, list)
            or any(not isinstance(sheet, str) or not sheet for sheet in sheets)
            or len(sheets) != len(set(sheets))
        ):
            raise ValueError(
                "xlsx_sheet_names_equal requires unique non-empty sheet names"
            )
    elif operation == "xlsx_table_names_equal":
        sheet = assertion.get("sheet")
        tables = assertion.get("tables")
        if not isinstance(sheet, str) or not sheet:
            raise ValueError("xlsx_table_names_equal requires a non-empty sheet")
        if (
            not isinstance(tables, list)
            or any(not isinstance(table, str) or not table for table in tables)
            or len(tables) != len(set(tables))
        ):
            raise ValueError(
                "xlsx_table_names_equal requires unique non-empty table names"
            )
    elif operation == "xlsx_table_properties":
        required_keys = {
            "op",
            "path",
            "sheet",
            "table_name",
            "data_range",
            "table_style",
        }
        if set(assertion) != required_keys:
            raise ValueError(
                "xlsx_table_properties requires exactly op, path, sheet, "
                "table_name, data_range, and table_style"
            )
        for field in ("sheet", "table_name", "table_style"):
            if not isinstance(assertion[field], str) or not assertion[field]:
                raise ValueError(
                    f"xlsx_table_properties.{field} must be a non-empty string"
                )
        _cell_range(
            assertion["data_range"],
            "xlsx_table_properties.data_range",
        )
    elif operation == "xlsx_table_definitions_equal":
        if set(assertion) != {"op", "path", "sheet", "tables"}:
            raise ValueError(
                "xlsx_table_definitions_equal requires exactly op, path, "
                "sheet, and tables"
            )
        if not isinstance(assertion["sheet"], str) or not assertion["sheet"]:
            raise ValueError(
                "xlsx_table_definitions_equal.sheet must be non-empty"
            )
        tables = assertion["tables"]
        if not isinstance(tables, list):
            raise ValueError(
                "xlsx_table_definitions_equal.tables must be a list"
            )
        canonical_tables = []
        for index, table in enumerate(tables):
            if (
                not isinstance(table, dict)
                or set(table) != {"data_range", "table_style"}
            ):
                raise ValueError(
                    "xlsx_table_definitions_equal table entries require "
                    "exactly data_range and table_style"
                )
            if not isinstance(table["table_style"], str) or not table["table_style"]:
                raise ValueError(
                    "xlsx_table_definitions_equal.table_style must be non-empty"
                )
            canonical_tables.append(
                (
                    _canonical_cell_range(
                        table["data_range"],
                        f"xlsx_table_definitions_equal.tables[{index}].data_range",
                    ),
                    table["table_style"],
                )
            )
        if len(canonical_tables) != len(set(canonical_tables)):
            raise ValueError(
                "xlsx_table_definitions_equal.tables must be unique"
            )
    elif operation == "xlsx_chart_count":
        count = assertion.get("count")
        if not isinstance(assertion.get("sheet"), str) or not assertion["sheet"]:
            raise ValueError("xlsx_chart_count requires a non-empty sheet")
        if not isinstance(count, int) or isinstance(count, bool) or count < 0:
            raise ValueError("xlsx_chart_count requires a non-negative integer count")
    elif operation == "xlsx_chart_properties":
        if set(assertion) != {"op", "path", "sheet", "index", "properties"}:
            raise ValueError(
                "xlsx_chart_properties requires exactly op, path, sheet, "
                "index, and properties"
            )
        if not isinstance(assertion["sheet"], str) or not assertion["sheet"]:
            raise ValueError("xlsx_chart_properties.sheet must be non-empty")
        index = assertion["index"]
        if not isinstance(index, int) or isinstance(index, bool) or index < 0:
            raise ValueError(
                "xlsx_chart_properties.index must be a non-negative 0-indexed integer"
            )
        properties = assertion["properties"]
        if (
            not isinstance(properties, dict)
            or set(properties) != XLSX_CHART_ASSERTION_PROPERTIES
        ):
            raise ValueError(
                "xlsx_chart_properties.properties requires exactly chart_type, "
                "target_cell, title, x_axis, y_axis, and data_range"
            )
        if properties["chart_type"] not in XLSX_CHART_TYPES:
            raise ValueError(
                "xlsx_chart_properties.properties.chart_type must be one of "
                f"{sorted(XLSX_CHART_TYPES)}"
            )
        _cell_coordinate(
            properties["target_cell"],
            "xlsx_chart_properties.properties.target_cell",
        )
        _cell_range(
            properties["data_range"],
            "xlsx_chart_properties.properties.data_range",
        )
        for field in ("title", "x_axis", "y_axis"):
            if properties[field] is not None and not isinstance(
                properties[field],
                str,
            ):
                raise ValueError(
                    f"xlsx_chart_properties.properties.{field} must be a string or null"
                )
    elif operation == "xlsx_cell_style":
        if not isinstance(assertion.get("sheet"), str) or not assertion["sheet"]:
            raise ValueError("xlsx_cell_style requires a non-empty sheet")
        _cell_coordinate(assertion.get("cell"), "xlsx_cell_style.cell")
        _validate_style_properties(
            assertion.get("properties"),
            "xlsx_cell_style.properties",
        )
    elif operation == "xlsx_merged_ranges_equal":
        if set(assertion) != {"op", "path", "sheet", "ranges"}:
            raise ValueError(
                "xlsx_merged_ranges_equal requires exactly op, path, sheet, and ranges"
            )
        if not isinstance(assertion["sheet"], str) or not assertion["sheet"]:
            raise ValueError("xlsx_merged_ranges_equal.sheet must be non-empty")
        ranges = assertion["ranges"]
        if not isinstance(ranges, list):
            raise ValueError("xlsx_merged_ranges_equal.ranges must be a list")
        canonical_ranges = [
            _canonical_cell_range(
                cell_range,
                "xlsx_merged_ranges_equal.ranges",
            )
            for cell_range in ranges
        ]
        if len(canonical_ranges) != len(set(canonical_ranges)):
            raise ValueError(
                "xlsx_merged_ranges_equal.ranges must contain unique ranges"
            )
    elif operation == "pdf_page_count":
        count = assertion.get("count")
        if not isinstance(count, int) or isinstance(count, bool) or count < 0:
            raise ValueError("pdf_page_count requires a non-negative integer count")
    elif operation in {"pdf_page_text_contains", "pdf_page_text_equals"}:
        page = assertion.get("page")
        if not isinstance(page, int) or isinstance(page, bool) or page < 1:
            raise ValueError(
                f"{operation} requires a positive 1-indexed page"
            )
        if not isinstance(assertion.get("value"), str):
            raise ValueError(f"{operation} requires a string value")


def _validate_recipe(recipe: Any, label: str) -> dict[str, Any]:
    if not isinstance(recipe, dict):
        raise ValueError(f"{label} must be an object")
    allowed = {"directories", "files", "remove"}
    unexpected = sorted(set(recipe) - allowed)
    if unexpected:
        raise ValueError(f"{label} has unsupported keys: {unexpected}")
    directories = recipe.get("directories", [])
    files = recipe.get("files", [])
    remove = recipe.get("remove", [])
    if not isinstance(directories, list) or not isinstance(files, list) or not isinstance(remove, list):
        raise ValueError(f"{label} directories/files/remove must be lists")
    if len(files) > 200:
        raise ValueError(f"{label} exceeds the 200-file safety limit")
    seen_paths = set()
    for directory in directories:
        normalized = safe_relative_path(directory).as_posix()
        if normalized in seen_paths:
            raise ValueError(f"{label} duplicates path {normalized!r}")
        seen_paths.add(normalized)
    for record in files:
        if not isinstance(record, dict):
            raise ValueError(f"{label} file records must be objects")
        if set(record) != {"path", "format", "content"}:
            raise ValueError(
                f"{label} file requires exactly path, format, and content"
            )
        normalized = safe_relative_path(record["path"]).as_posix()
        if normalized in seen_paths:
            raise ValueError(f"{label} duplicates path {normalized!r}")
        seen_paths.add(normalized)
        if record["format"] not in SUPPORTED_FILE_FORMATS:
            raise ValueError(f"unsupported generated file format: {record['format']!r}")
        if record["format"] == "xlsx":
            _validate_xlsx_content(record["content"])
    for path in remove:
        safe_relative_path(path)
    return recipe


def validate_task_spec_structure(task: dict[str, Any]) -> None:
    _validate_recipe(task.get("initial_workspace"), "initial_workspace")
    _validate_recipe(task.get("oracle_final_state"), "oracle_final_state")
    evaluator = task.get("evaluator")
    if not isinstance(evaluator, dict):
        raise ValueError("evaluator must be an object")
    if evaluator.get("type") != "workspace_assertions_v1":
        raise ValueError("only declarative workspace_assertions_v1 evaluators are allowed")
    assertions = evaluator.get("assertions")
    if not isinstance(assertions, list) or not assertions:
        raise ValueError("evaluator.assertions must be a non-empty list")
    for assertion in assertions:
        if not isinstance(assertion, dict):
            raise ValueError("evaluator assertions must be objects")
        operation = assertion.get("op")
        if operation not in SUPPORTED_ASSERTIONS:
            raise ValueError(f"unsupported evaluator assertion: {operation!r}")
        _validate_assertion_arguments(assertion)
        if operation == "tree_equals":
            paths = assertion.get("paths")
            if not isinstance(paths, list):
                raise ValueError("tree_equals requires a paths list")
            for path in paths:
                safe_relative_path(path)
        else:
            safe_relative_path(assertion.get("path"))


def _write_generated_file(path: Path, file_format: str, content: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if file_format == "text":
        if not isinstance(content, str):
            raise ValueError("text file content must be a string")
        path.write_text(content, encoding="utf-8")
    elif file_format == "json":
        path.write_text(
            json.dumps(content, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
    elif file_format == "csv":
        if not isinstance(content, list) or any(not isinstance(row, list) for row in content):
            raise ValueError("csv content must be a list of row lists")
        with path.open("w", encoding="utf-8", newline="") as handle:
            csv.writer(handle).writerows(content)
    elif file_format == "xlsx":
        try:
            from openpyxl import Workbook
            from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, Reference
            from openpyxl.styles import Alignment, PatternFill
            from openpyxl.worksheet.table import Table, TableStyleInfo
        except ImportError as exc:
            raise RuntimeError("xlsx generation requires openpyxl") from exc
        _validate_xlsx_content(content)
        workbook = Workbook()
        workbook.remove(workbook.active)
        for sheet_spec in content["sheets"]:
            sheet = workbook.create_sheet(title=sheet_spec["name"])
            for row in sheet_spec["rows"]:
                sheet.append(row)
            for style_spec in sheet_spec.get("styles", []):
                start_column, start_row = _cell_coordinate(
                    style_spec["start_cell"],
                    "xlsx style start_cell",
                )
                end_column, end_row = _cell_coordinate(
                    style_spec["end_cell"],
                    "xlsx style end_cell",
                )
                for cells in sheet.iter_rows(
                    min_col=start_column,
                    min_row=start_row,
                    max_col=end_column,
                    max_row=end_row,
                ):
                    for cell in cells:
                        if "bold" in style_spec or "font_color" in style_spec:
                            font = copy(cell.font)
                            if "bold" in style_spec:
                                font.bold = style_spec["bold"]
                            if "font_color" in style_spec:
                                font.color = (
                                    "FF"
                                    + _normalize_hex_color(
                                        style_spec["font_color"],
                                        "font_color",
                                    )
                                )
                            cell.font = font
                        if "bg_color" in style_spec:
                            cell.fill = PatternFill(
                                fill_type="solid",
                                fgColor=(
                                    "FF"
                                    + _normalize_hex_color(
                                        style_spec["bg_color"],
                                        "bg_color",
                                    )
                                ),
                            )
                        if "alignment" in style_spec:
                            alignment = copy(cell.alignment)
                            alignment.horizontal = style_spec["alignment"]
                            cell.alignment = alignment
            for merged_range in sheet_spec.get("merged_ranges", []):
                sheet.merge_cells(merged_range)
            for table_spec in sheet_spec.get("tables", []):
                table = Table(
                    displayName=table_spec["table_name"],
                    ref=table_spec["data_range"],
                )
                table.tableStyleInfo = TableStyleInfo(
                    name=table_spec["table_style"],
                    showFirstColumn=False,
                    showLastColumn=False,
                    showRowStripes=True,
                    showColumnStripes=False,
                )
                sheet.add_table(table)
            for chart_spec in sheet_spec.get("charts", []):
                chart_type = chart_spec["chart_type"]
                if chart_type == "bar":
                    chart = BarChart()
                elif chart_type == "line":
                    chart = LineChart()
                elif chart_type == "pie":
                    chart = PieChart()
                elif chart_type == "area":
                    chart = AreaChart()
                else:
                    raise AssertionError(f"unhandled chart type: {chart_type}")
                min_column, min_row, max_column, max_row = _cell_range(
                    chart_spec["data_range"],
                    "chart data_range",
                )
                data_min_column = (
                    min_column + 1 if max_column > min_column else min_column
                )
                data = Reference(
                    sheet,
                    min_col=data_min_column,
                    min_row=min_row,
                    max_col=max_column,
                    max_row=max_row,
                )
                chart.add_data(data, titles_from_data=max_row > min_row)
                if max_column > min_column and max_row > min_row:
                    categories = Reference(
                        sheet,
                        min_col=min_column,
                        min_row=min_row + 1,
                        max_row=max_row,
                    )
                    chart.set_categories(categories)
                if "title" in chart_spec:
                    chart.title = chart_spec["title"]
                if "x_axis" in chart_spec and hasattr(chart, "x_axis"):
                    chart.x_axis.title = chart_spec["x_axis"]
                if "y_axis" in chart_spec and hasattr(chart, "y_axis"):
                    chart.y_axis.title = chart_spec["y_axis"]
                sheet.add_chart(chart, chart_spec["target_cell"])
        workbook.properties.created = DETERMINISTIC_OFFICE_TIMESTAMP
        workbook.properties.modified = DETERMINISTIC_OFFICE_TIMESTAMP
        workbook.save(path)
        normalized_path = path.with_name(path.name + ".normalized")
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
            normalized_path,
            "w",
        ) as destination:
            for source_info in source.infolist():
                member_data = source.read(source_info.filename)
                if source_info.filename == "docProps/core.xml":
                    member_data = re.sub(
                        rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                        rb"\g<1>2000-01-01T00:00:00Z\g<2>",
                        member_data,
                    )
                target_info = zipfile.ZipInfo(
                    source_info.filename,
                    date_time=DETERMINISTIC_ZIP_TIMESTAMP,
                )
                target_info.compress_type = source_info.compress_type
                target_info.external_attr = source_info.external_attr
                target_info.create_system = source_info.create_system
                destination.writestr(target_info, member_data)
        normalized_path.replace(path)
    elif file_format == "pdf":
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.pdfgen import canvas
        except ImportError as exc:
            raise RuntimeError("pdf generation requires reportlab") from exc
        if not isinstance(content, list) or any(not isinstance(page, str) for page in content):
            raise ValueError("pdf content must be a list of page strings")
        document = canvas.Canvas(
            str(path),
            pagesize=letter,
            invariant=1,
        )
        for page in content:
            text = document.beginText(54, 738)
            for line in page.splitlines() or [""]:
                text.textLine(line)
            document.drawText(text)
            document.showPage()
        document.save()
    elif file_format == "base64":
        if not isinstance(content, str):
            raise ValueError("base64 file content must be a string")
        path.write_bytes(base64.b64decode(content, validate=True))
    else:
        raise AssertionError(f"unhandled file format: {file_format}")
    if path.stat().st_size > 10 * 1024 * 1024:
        raise ValueError(f"generated file exceeds 10 MiB: {path.name}")


def apply_workspace_recipe(
    workspace: str | Path,
    recipe: dict[str, Any],
    *,
    require_empty: bool = False,
) -> None:
    recipe = _validate_recipe(recipe, "workspace recipe")
    root = Path(workspace).resolve()
    root.mkdir(parents=True, exist_ok=True)
    if require_empty and any(root.iterdir()):
        raise ValueError(f"workspace must be empty before initialization: {root}")
    for relative in recipe.get("remove", []):
        target = root / safe_relative_path(relative)
        if target.is_dir() and not target.is_symlink():
            shutil.rmtree(target)
        elif target.exists() or target.is_symlink():
            target.unlink()
    for relative in recipe.get("directories", []):
        (root / safe_relative_path(relative)).mkdir(parents=True, exist_ok=True)
    for record in recipe.get("files", []):
        _write_generated_file(
            root / safe_relative_path(record["path"]),
            record["format"],
            record["content"],
        )


def _read_csv(path: Path) -> list[list[str]]:
    with path.open("r", encoding="utf-8", newline="") as handle:
        return list(csv.reader(handle))


def _chart_text(title: Any) -> str | None:
    if title is None:
        return None
    text = getattr(title, "tx", None)
    rich_text = getattr(text, "rich", None)
    paragraphs = getattr(rich_text, "p", None)
    if paragraphs is not None:
        rendered_paragraphs = []
        for paragraph in paragraphs:
            fragments = [
                run.t
                for run in (getattr(paragraph, "r", None) or [])
                if isinstance(getattr(run, "t", None), str)
            ]
            fragments.extend(
                field.t
                for field in (getattr(paragraph, "fld", None) or [])
                if isinstance(getattr(field, "t", None), str)
            )
            rendered_paragraphs.append("".join(fragments))
        return "\n".join(rendered_paragraphs)
    string_reference = getattr(text, "strRef", None)
    string_cache = getattr(string_reference, "strCache", None)
    points = getattr(string_cache, "pt", None)
    if points is not None:
        return "".join(
            point.v
            for point in points
            if isinstance(getattr(point, "v", None), str)
        )
    return None


def _chart_formula_references(chart: Any) -> list[str]:
    references: list[str] = []
    for series in chart.ser:
        sources = [getattr(series, "tx", None)]
        sources.extend(
            getattr(series, attribute, None)
            for attribute in ("val", "cat", "xVal", "yVal")
        )
        for source in sources:
            if source is None:
                continue
            for reference_name in ("strRef", "numRef", "multiLvlStrRef"):
                reference = getattr(source, reference_name, None)
                formula = getattr(reference, "f", None)
                if isinstance(formula, str) and formula not in references:
                    references.append(formula)
    return references


def _formula_reference_bounds(
    formula: str,
    *,
    expected_sheet: str,
) -> tuple[int, int, int, int] | None:
    formula = formula.removeprefix("=")
    if "!" not in formula:
        return None
    sheet_reference, address = formula.rsplit("!", 1)
    if sheet_reference.startswith("'") and sheet_reference.endswith("'"):
        sheet_reference = sheet_reference[1:-1].replace("''", "'")
    if sheet_reference != expected_sheet:
        return None
    address = address.replace("$", "")
    try:
        if ":" in address:
            return _cell_range(address, "chart series reference")
        column, row = _cell_coordinate(address, "chart series reference")
        return column, row, column, row
    except ValueError:
        return None


def _chart_data_range(
    chart: Any,
    *,
    sheet_name: str,
) -> tuple[str | None, list[str]]:
    references = _chart_formula_references(chart)
    bounds = [
        parsed
        for formula in references
        if (
            parsed := _formula_reference_bounds(
                formula,
                expected_sheet=sheet_name,
            )
        )
        is not None
    ]
    if not bounds:
        return None, references
    min_column = min(bound[0] for bound in bounds)
    min_row = min(bound[1] for bound in bounds)
    max_column = max(bound[2] for bound in bounds)
    max_row = max(bound[3] for bound in bounds)
    # Toolathlon's Excel chart tool interprets a rectangular input as one
    # category column followed by one or more value-series columns, with the
    # first row holding series names. A bounding box alone is insufficient:
    # deleting a middle series can leave the same outer rectangle. Require the
    # complete reference set so missing or extra series fail the assertion.
    if (
        min_column >= max_column
        or min_row >= max_row
        or len(chart.ser) != max_column - min_column
    ):
        return None, references
    expected_bounds = {
        (min_column, min_row + 1, min_column, max_row),
    }
    for column in range(min_column + 1, max_column + 1):
        expected_bounds.add((column, min_row, column, min_row))
        expected_bounds.add((column, min_row + 1, column, max_row))
    if set(bounds) != expected_bounds:
        return None, references
    return (
        f"{_column_name(min_column)}{min_row}:"
        f"{_column_name(max_column)}{max_row}",
        references,
    )


def _chart_type(chart: Any) -> str | None:
    class_name = type(chart).__name__
    if class_name.startswith("BarChart"):
        # The Excel MCP API calls this family "bar", while openpyxl stores its
        # default OOXML orientation as ``type="col"``. The class family is the
        # stable representation of the public tool argument.
        return "bar"
    if class_name.startswith("LineChart"):
        return "line"
    if class_name.startswith("PieChart"):
        return "pie"
    if class_name.startswith("AreaChart"):
        return "area"
    return None


def _chart_anchor(chart: Any) -> tuple[str | None, dict[str, Any]]:
    anchor = chart.anchor
    if isinstance(anchor, str):
        return (
            _canonical_cell_coordinate(anchor, "chart anchor"),
            {"anchor_type": "cell", "coordinate": anchor},
        )
    marker = getattr(anchor, "_from", None)
    evidence = {"anchor_type": type(anchor).__name__}
    if marker is None:
        return None, evidence
    target_cell = f"{_column_name(marker.col + 1)}{marker.row + 1}"
    evidence.update(
        {
            "target_cell": target_cell,
            "column_offset": marker.colOff,
            "row_offset": marker.rowOff,
        }
    )
    return target_cell, evidence


def _chart_properties(
    chart: Any,
    *,
    sheet_name: str,
) -> tuple[dict[str, Any], dict[str, Any]]:
    data_range, references = _chart_data_range(chart, sheet_name=sheet_name)
    target_cell, anchor_evidence = _chart_anchor(chart)
    properties = {
        "chart_type": _chart_type(chart),
        "target_cell": target_cell,
        "title": _chart_text(chart.title),
        "x_axis": _chart_text(getattr(getattr(chart, "x_axis", None), "title", None)),
        "y_axis": _chart_text(getattr(getattr(chart, "y_axis", None), "title", None)),
        "data_range": data_range,
    }
    return properties, {
        "data_range_recovered_from_series_references": references,
        "anchor": anchor_evidence,
    }


def _evaluate_assertion(root: Path, assertion: dict[str, Any]) -> tuple[bool, str]:
    operation = assertion["op"]
    if operation == "tree_equals":
        expected = sorted(safe_relative_path(path).as_posix() for path in assertion["paths"])
        actual = sorted(
            path.relative_to(root).as_posix()
            for path in root.rglob("*")
            if not path.is_symlink()
        )
        return actual == expected, f"tree_equals expected={expected} actual={actual}"

    path = root / safe_relative_path(assertion["path"])
    if operation == "exists":
        kind = assertion.get("kind", "any")
        passed = path.exists() and (
            kind == "any"
            or (kind == "file" and path.is_file())
            or (kind == "directory" and path.is_dir())
        )
        return passed, f"exists path={assertion['path']} kind={kind}"
    if operation == "absent":
        return not path.exists(), f"absent path={assertion['path']}"
    if not path.is_file():
        return False, f"missing file path={assertion['path']}"

    if operation == "text_equals":
        actual = path.read_text(encoding="utf-8")
        return actual == assertion.get("value"), f"text_equals path={assertion['path']}"
    if operation == "text_contains":
        actual = path.read_text(encoding="utf-8")
        return str(assertion.get("value")) in actual, f"text_contains path={assertion['path']}"
    if operation == "json_equals":
        actual = json.loads(path.read_text(encoding="utf-8"))
        return actual == assertion.get("value"), f"json_equals path={assertion['path']}"
    if operation == "csv_equals":
        expected_rows = [[str(cell) for cell in row] for row in assertion.get("rows", [])]
        return _read_csv(path) == expected_rows, f"csv_equals path={assertion['path']}"
    if operation in {"xlsx_cells_equal", "xlsx_nonempty_cells_equal"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("xlsx evaluation requires openpyxl") from exc
        workbook = load_workbook(path, data_only=False, read_only=True)
        sheet_name = assertion.get("sheet")
        cells = assertion.get("cells")
        if not isinstance(sheet_name, str) or not isinstance(cells, dict):
            raise ValueError("xlsx_cells_equal requires sheet and cells")
        if sheet_name not in workbook.sheetnames:
            workbook.close()
            return False, f"xlsx missing sheet={sheet_name}"
        sheet = workbook[sheet_name]
        expected = {
            _canonical_cell_coordinate(coordinate, f"{operation}.cells"): value
            for coordinate, value in cells.items()
            if operation == "xlsx_cells_equal" or value is not None
        }
        if operation == "xlsx_cells_equal":
            actual = {
                coordinate: sheet[coordinate].value
                for coordinate in expected
            }
        else:
            actual = {
                cell.coordinate: cell.value
                for row in sheet.iter_rows()
                for cell in row
                if cell.value is not None
            }
        result = (
            actual == expected,
            f"{operation} path={assertion['path']} sheet={sheet_name} "
            f"expected={expected} actual={actual}",
        )
        workbook.close()
        return result
    if operation in {
        "xlsx_sheet_names_equal",
        "xlsx_table_names_equal",
        "xlsx_table_properties",
        "xlsx_table_definitions_equal",
        "xlsx_chart_count",
        "xlsx_chart_properties",
        "xlsx_cell_style",
        "xlsx_merged_ranges_equal",
    }:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("xlsx evaluation requires openpyxl") from exc
        workbook = load_workbook(path, data_only=False, read_only=False)
        try:
            if operation == "xlsx_sheet_names_equal":
                expected_sheets = assertion["sheets"]
                actual_sheets = workbook.sheetnames
                return (
                    actual_sheets == expected_sheets,
                    f"xlsx_sheet_names_equal path={assertion['path']} "
                    f"expected={expected_sheets} actual={actual_sheets}",
                )
            sheet_name = assertion["sheet"]
            if sheet_name not in workbook.sheetnames:
                return False, f"xlsx missing sheet={sheet_name}"
            sheet = workbook[sheet_name]
            if operation == "xlsx_table_names_equal":
                expected_tables = sorted(assertion["tables"])
                actual_tables = sorted(sheet.tables.keys())
                return (
                    actual_tables == expected_tables,
                    f"xlsx_table_names_equal path={assertion['path']} "
                    f"sheet={sheet_name} expected={expected_tables} "
                    f"actual={actual_tables}",
                )
            if operation == "xlsx_table_properties":
                table_name = assertion["table_name"]
                if table_name not in sheet.tables:
                    return (
                        False,
                        f"xlsx_table_properties path={assertion['path']} "
                        f"sheet={sheet_name} missing_table={table_name}",
                    )
                table = sheet.tables[table_name]
                expected_properties = {
                    "data_range": _canonical_cell_range(
                        assertion["data_range"],
                        "xlsx_table_properties.data_range",
                    ),
                    "table_style": assertion["table_style"],
                }
                actual_properties = {
                    "data_range": _canonical_cell_range(
                        table.ref,
                        "actual table data_range",
                    ),
                    "table_style": (
                        table.tableStyleInfo.name
                        if table.tableStyleInfo is not None
                        else None
                    ),
                }
                return (
                    actual_properties == expected_properties,
                    f"xlsx_table_properties path={assertion['path']} "
                    f"sheet={sheet_name} table={table_name} "
                    f"expected={expected_properties} actual={actual_properties}",
                )
            if operation == "xlsx_table_definitions_equal":
                expected_tables = sorted(
                    (
                        _canonical_cell_range(
                            table["data_range"],
                            "xlsx_table_definitions_equal.data_range",
                        ),
                        table["table_style"],
                    )
                    for table in assertion["tables"]
                )
                actual_tables = sorted(
                    (
                        _canonical_cell_range(
                            table.ref,
                            "actual table data_range",
                        ),
                        (
                            table.tableStyleInfo.name
                            if table.tableStyleInfo is not None
                            else None
                        ),
                    )
                    for table in sheet.tables.values()
                )
                return (
                    actual_tables == expected_tables,
                    f"xlsx_table_definitions_equal path={assertion['path']} "
                    f"sheet={sheet_name} expected={expected_tables} "
                    f"actual={actual_tables}",
                )
            if operation == "xlsx_chart_count":
                expected_count = assertion["count"]
                actual_count = len(sheet._charts)
                return (
                    actual_count == expected_count,
                    f"xlsx_chart_count path={assertion['path']} "
                    f"sheet={sheet_name} expected={expected_count} "
                    f"actual={actual_count}",
                )
            if operation == "xlsx_chart_properties":
                chart_index = assertion["index"]
                if chart_index >= len(sheet._charts):
                    return (
                        False,
                        f"xlsx_chart_properties path={assertion['path']} "
                        f"sheet={sheet_name} missing_0_indexed_chart={chart_index} "
                        f"chart_count={len(sheet._charts)}",
                    )
                expected_properties = dict(assertion["properties"])
                expected_properties["target_cell"] = _canonical_cell_coordinate(
                    expected_properties["target_cell"],
                    "xlsx_chart_properties.properties.target_cell",
                )
                expected_properties["data_range"] = _canonical_cell_range(
                    expected_properties["data_range"],
                    "xlsx_chart_properties.properties.data_range",
                )
                actual_properties, recovery_evidence = _chart_properties(
                    sheet._charts[chart_index],
                    sheet_name=sheet_name,
                )
                return (
                    actual_properties == expected_properties,
                    f"xlsx_chart_properties path={assertion['path']} "
                    f"sheet={sheet_name} index={chart_index} "
                    f"expected={expected_properties} actual={actual_properties} "
                    f"stable_recovery={recovery_evidence}",
                )
            if operation == "xlsx_merged_ranges_equal":
                expected_ranges = sorted(
                    _canonical_cell_range(
                        cell_range,
                        "xlsx_merged_ranges_equal.ranges",
                    )
                    for cell_range in assertion["ranges"]
                )
                actual_ranges = sorted(
                    _canonical_cell_range(
                        str(cell_range),
                        "actual merged cell range",
                    )
                    for cell_range in sheet.merged_cells.ranges
                )
                return (
                    actual_ranges == expected_ranges,
                    f"xlsx_merged_ranges_equal path={assertion['path']} "
                    f"sheet={sheet_name} expected={expected_ranges} "
                    f"actual={actual_ranges}",
                )
            cell = sheet[assertion["cell"]]
            expected_properties = assertion["properties"]
            actual_properties: dict[str, Any] = {}
            for property_name in expected_properties:
                if property_name == "bold":
                    actual_properties[property_name] = bool(cell.font.bold)
                elif property_name == "font_color":
                    color = cell.font.color
                    actual_properties[property_name] = (
                        _normalize_hex_color(color.rgb, "actual font color")
                        if color is not None
                        and color.type == "rgb"
                        and isinstance(color.rgb, str)
                        else None
                    )
                elif property_name == "bg_color":
                    color = cell.fill.fgColor
                    actual_properties[property_name] = (
                        _normalize_hex_color(color.rgb, "actual background color")
                        if color is not None
                        and cell.fill.patternType == "solid"
                        and color.type == "rgb"
                        and isinstance(color.rgb, str)
                        else None
                    )
                elif property_name == "alignment":
                    actual_properties[property_name] = cell.alignment.horizontal
                else:
                    raise AssertionError(
                        f"unhandled xlsx style property: {property_name}"
                    )
            normalized_expected = dict(expected_properties)
            for color_property in ("font_color", "bg_color"):
                if color_property in normalized_expected:
                    normalized_expected[color_property] = _normalize_hex_color(
                        normalized_expected[color_property],
                        f"expected {color_property}",
                    )
            return (
                actual_properties == normalized_expected,
                f"xlsx_cell_style path={assertion['path']} sheet={sheet_name} "
                f"cell={assertion['cell']} expected={normalized_expected} "
                f"actual={actual_properties}",
            )
        finally:
            workbook.close()
    if operation == "pdf_text_contains":
        try:
            import fitz
        except ImportError as exc:
            raise RuntimeError("PDF evaluation requires PyMuPDF") from exc
        with fitz.open(path) as document:
            actual = "\n".join(page.get_text() for page in document)
        return str(assertion.get("value")) in actual, f"pdf_text_contains path={assertion['path']}"
    if operation in {
        "pdf_page_count",
        "pdf_page_text_contains",
        "pdf_page_text_equals",
    }:
        try:
            import fitz
        except ImportError as exc:
            raise RuntimeError("PDF evaluation requires PyMuPDF") from exc
        with fitz.open(path) as document:
            if operation == "pdf_page_count":
                expected_count = assertion["count"]
                actual_count = document.page_count
                return (
                    actual_count == expected_count,
                    f"pdf_page_count path={assertion['path']} "
                    f"expected={expected_count} actual={actual_count}",
                )
            page_number = assertion["page"]
            if page_number > document.page_count:
                return (
                    False,
                    f"{operation} path={assertion['path']} "
                    f"missing_page={page_number} page_count={document.page_count}",
                )
            actual = document.load_page(page_number - 1).get_text()
            expected_value = assertion["value"]
            if operation == "pdf_page_text_equals":
                normalized_actual = " ".join(actual.split())
                normalized_expected = " ".join(expected_value.split())
                return (
                    normalized_actual == normalized_expected,
                    f"pdf_page_text_equals path={assertion['path']} "
                    f"page={page_number} expected={normalized_expected!r} "
                    f"actual={normalized_actual!r}",
                )
            return (
                expected_value in actual,
                f"pdf_page_text_contains path={assertion['path']} "
                f"page={page_number} value={expected_value!r}",
            )
    if operation == "sha256":
        actual = hashlib.sha256(path.read_bytes()).hexdigest()
        return actual == assertion.get("value"), f"sha256 path={assertion['path']}"
    if operation == "file_size_at_least":
        minimum = assertion.get("bytes")
        if not isinstance(minimum, int) or minimum < 0:
            raise ValueError("file_size_at_least requires non-negative integer bytes")
        return path.stat().st_size >= minimum, f"file_size_at_least path={assertion['path']}"
    raise AssertionError(f"unhandled evaluator assertion: {operation}")


def evaluate_workspace(
    workspace: str | Path,
    evaluator: dict[str, Any],
) -> dict[str, Any]:
    validate_task_spec_structure(
        {
            "initial_workspace": {},
            "oracle_final_state": {},
            "evaluator": evaluator,
        }
    )
    root = Path(workspace).resolve()
    checks = []
    for index, assertion in enumerate(evaluator["assertions"]):
        try:
            passed, detail = _evaluate_assertion(root, assertion)
        except Exception as exc:
            passed = False
            detail = f"{type(exc).__name__}: {exc}"
        checks.append({"index": index, "passed": passed, "detail": detail})
    return {
        "passed": all(check["passed"] for check in checks),
        "version": "workspace_assertions_v1",
        "checks": checks,
    }


def workspace_digest(
    workspace: str | Path,
    *,
    allow_symlinks: bool = False,
) -> str:
    root = Path(workspace).resolve()
    entries = []
    for path in sorted(root.rglob("*")):
        relative = path.relative_to(root).as_posix()
        if path.is_symlink():
            if not allow_symlinks:
                raise ValueError(
                    f"synthetic workspace cannot contain symlinks: {relative}"
                )
            entries.append(
                {
                    "path": relative,
                    "kind": "symlink",
                    "target": str(path.readlink()),
                }
            )
            continue
        if path.is_dir():
            entries.append({"path": relative, "kind": "directory"})
        else:
            entries.append(
                {
                    "path": relative,
                    "kind": "file",
                    "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                }
            )
    return hashlib.sha256(
        json.dumps(entries, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()


def verify_task_assets(task: dict[str, Any]) -> dict[str, Any]:
    validate_task_spec_structure(task)
    with tempfile.TemporaryDirectory(prefix="tokmem_toolathlon_verify_") as temporary:
        initial = Path(temporary) / "initial"
        oracle = Path(temporary) / "oracle"
        apply_workspace_recipe(initial, task["initial_workspace"], require_empty=True)
        shutil.copytree(initial, oracle)
        initial_result = evaluate_workspace(initial, task["evaluator"])
        apply_workspace_recipe(oracle, task["oracle_final_state"])
        oracle_result = evaluate_workspace(oracle, task["evaluator"])
        passed = not initial_result["passed"] and oracle_result["passed"]
        reasons = []
        if initial_result["passed"]:
            reasons.append("evaluator already passes on the untouched initial workspace")
        if not oracle_result["passed"]:
            reasons.append("evaluator fails on the declared oracle final state")
        return {
            "passed": passed,
            "reasons": reasons,
            "initial_evaluator": initial_result,
            "oracle_evaluator": oracle_result,
            "initial_workspace_hash": workspace_digest(initial),
            "oracle_workspace_hash": workspace_digest(oracle),
        }
