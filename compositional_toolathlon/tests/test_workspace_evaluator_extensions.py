from __future__ import annotations

import importlib.util
import tempfile
import unittest
from pathlib import Path

from compositional_toolathlon.synthetic_workspace import (
    apply_workspace_recipe,
    evaluate_workspace,
    validate_task_spec_structure,
)


OPENPYXL_AVAILABLE = importlib.util.find_spec("openpyxl") is not None
PDF_RUNTIME_AVAILABLE = (
    importlib.util.find_spec("reportlab") is not None
    and importlib.util.find_spec("fitz") is not None
)


def _evaluator(*assertions):
    return {
        "type": "workspace_assertions_v1",
        "assertions": list(assertions),
    }


@unittest.skipUnless(
    OPENPYXL_AVAILABLE,
    "openpyxl is installed only in the Toolathlon runtime",
)
class XlsxEvaluatorExtensionTests(unittest.TestCase):
    def make_recipe(self):
        return {
            "directories": ["reports"],
            "files": [
                {
                    "path": "reports/sales.xlsx",
                    "format": "xlsx",
                    "content": {
                        "sheets": [
                            {"name": "Cover", "rows": [["Quarterly sales"]]},
                            {
                                "name": "Data",
                                "rows": [
                                    ["Region", "Units", "Revenue"],
                                    ["North", 2, 20],
                                    ["South", 3, 30],
                                ],
                                "styles": [
                                    {
                                        "start_cell": "A1",
                                        "end_cell": "C1",
                                        "bold": True,
                                        "font_color": "FFFFFF",
                                        "bg_color": "#4472C4",
                                        "alignment": "center",
                                    }
                                ],
                                "tables": [
                                    {
                                        "data_range": "A1:C3",
                                        "table_name": "SalesTable",
                                        "table_style": "TableStyleMedium9",
                                    }
                                ],
                                "charts": [
                                    {
                                        "data_range": "A1:C3",
                                        "chart_type": "bar",
                                        "target_cell": "E2",
                                        "title": "Sales overview",
                                        "x_axis": "Region",
                                        "y_axis": "Value",
                                    }
                                ],
                            },
                        ]
                    },
                }
            ],
            "remove": [],
        }

    def test_recipe_constructs_all_declared_xlsx_states(self):
        assertions = _evaluator(
            {
                "op": "xlsx_sheet_names_equal",
                "path": "reports/sales.xlsx",
                "sheets": ["Cover", "Data"],
            },
            {
                "op": "xlsx_table_names_equal",
                "path": "reports/sales.xlsx",
                "sheet": "Data",
                "tables": ["SalesTable"],
            },
            {
                "op": "xlsx_chart_count",
                "path": "reports/sales.xlsx",
                "sheet": "Data",
                "count": 1,
            },
            {
                "op": "xlsx_cell_style",
                "path": "reports/sales.xlsx",
                "sheet": "Data",
                "cell": "B1",
                "properties": {
                    "bold": True,
                    "font_color": "#FFFFFF",
                    "bg_color": "4472C4",
                    "alignment": "center",
                },
            },
        )
        with tempfile.TemporaryDirectory() as temporary:
            apply_workspace_recipe(
                temporary,
                self.make_recipe(),
                require_empty=True,
            )
            result = evaluate_workspace(temporary, assertions)
        self.assertTrue(result["passed"], result)
        self.assertTrue(all(check["passed"] for check in result["checks"]))

    def test_xlsx_assertions_reject_incorrect_final_state(self):
        assertions = _evaluator(
            {
                "op": "xlsx_sheet_names_equal",
                "path": "reports/sales.xlsx",
                "sheets": ["Data", "Cover"],
            },
            {
                "op": "xlsx_table_names_equal",
                "path": "reports/sales.xlsx",
                "sheet": "Data",
                "tables": [],
            },
            {
                "op": "xlsx_chart_count",
                "path": "reports/sales.xlsx",
                "sheet": "Data",
                "count": 2,
            },
            {
                "op": "xlsx_cell_style",
                "path": "reports/sales.xlsx",
                "sheet": "Data",
                "cell": "A1",
                "properties": {"bold": False},
            },
        )
        with tempfile.TemporaryDirectory() as temporary:
            apply_workspace_recipe(
                temporary,
                self.make_recipe(),
                require_empty=True,
            )
            result = evaluate_workspace(temporary, assertions)
        self.assertFalse(result["passed"])
        self.assertTrue(all(not check["passed"] for check in result["checks"]))

    def test_xlsx_recipe_rejects_nonlocal_ranges(self):
        recipe = self.make_recipe()
        recipe["files"][0]["content"]["sheets"][1]["tables"][0]["data_range"] = (
            "Other!A1:C3"
        )
        with tempfile.TemporaryDirectory() as temporary:
            with self.assertRaisesRegex(ValueError, "local A1-style cell"):
                apply_workspace_recipe(temporary, recipe, require_empty=True)

    def _add_declared_merge(self, workspace):
        from openpyxl import load_workbook

        path = Path(workspace) / "reports/sales.xlsx"
        workbook = load_workbook(path)
        workbook["Data"].merge_cells("A5:C5")
        workbook.save(path)
        workbook.close()

    def test_recipe_constructs_declared_merged_ranges(self):
        recipe = self.make_recipe()
        recipe["files"][0]["content"]["sheets"][1]["merged_ranges"] = ["A5:C5"]
        evaluator = _evaluator(
            {
                "op": "xlsx_merged_ranges_equal",
                "path": "reports/sales.xlsx",
                "sheet": "Data",
                "ranges": ["A5:C5"],
            }
        )
        with tempfile.TemporaryDirectory() as temporary:
            apply_workspace_recipe(temporary, recipe, require_empty=True)
            result = evaluate_workspace(temporary, evaluator)
        self.assertTrue(result["passed"], result)

    def test_recipe_rejects_overlapping_merged_ranges(self):
        recipe = self.make_recipe()
        recipe["files"][0]["content"]["sheets"][1]["merged_ranges"] = [
            "A5:B5",
            "B5:C5",
        ]
        with tempfile.TemporaryDirectory() as temporary:
            with self.assertRaisesRegex(ValueError, "must not overlap"):
                apply_workspace_recipe(temporary, recipe, require_empty=True)

    def chart_properties(self):
        return {
            "chart_type": "bar",
            "target_cell": "E2",
            "title": "Sales overview",
            "x_axis": "Region",
            "y_axis": "Value",
            "data_range": "A1:C3",
        }

    def test_strict_table_chart_and_merged_range_properties_pass(self):
        evaluator = _evaluator(
            {
                "op": "xlsx_table_properties",
                "path": "reports/sales.xlsx",
                "sheet": "Data",
                "table_name": "SalesTable",
                "data_range": "A1:C3",
                "table_style": "TableStyleMedium9",
            },
            {
                "op": "xlsx_chart_properties",
                "path": "reports/sales.xlsx",
                "sheet": "Data",
                "index": 0,
                "properties": self.chart_properties(),
            },
            {
                "op": "xlsx_merged_ranges_equal",
                "path": "reports/sales.xlsx",
                "sheet": "Data",
                "ranges": ["A5:C5"],
            },
        )
        with tempfile.TemporaryDirectory() as temporary:
            apply_workspace_recipe(
                temporary,
                self.make_recipe(),
                require_empty=True,
            )
            self._add_declared_merge(temporary)
            result = evaluate_workspace(temporary, evaluator)
        self.assertTrue(result["passed"], result)
        self.assertTrue(all(check["passed"] for check in result["checks"]))
        chart_detail = result["checks"][1]["detail"]
        self.assertIn("data_range_recovered_from_series_references", chart_detail)
        self.assertIn("OneCellAnchor", chart_detail)

    def test_nonempty_cell_set_rejects_extra_values(self):
        expected_cells = {
            "A1": "Region",
            "B1": "Units",
            "C1": "Revenue",
            "A2": "North",
            "B2": 2,
            "C2": 20,
            "A3": "South",
            "B3": 3,
            "C3": 30,
        }
        assertion = {
            "op": "xlsx_nonempty_cells_equal",
            "path": "reports/sales.xlsx",
            "sheet": "Data",
            "cells": expected_cells,
        }
        with tempfile.TemporaryDirectory() as temporary:
            apply_workspace_recipe(
                temporary,
                self.make_recipe(),
                require_empty=True,
            )
            clean = evaluate_workspace(temporary, _evaluator(assertion))

            from openpyxl import load_workbook

            path = Path(temporary) / "reports/sales.xlsx"
            workbook = load_workbook(path)
            workbook["Data"]["Z100"] = "EXTRA"
            workbook.save(path)
            workbook.close()
            tampered = evaluate_workspace(temporary, _evaluator(assertion))
        self.assertTrue(clean["passed"], clean)
        self.assertFalse(tampered["passed"], tampered)

    def test_table_definitions_can_ignore_random_names_but_enforce_structure(self):
        assertion = {
            "op": "xlsx_table_definitions_equal",
            "path": "reports/sales.xlsx",
            "sheet": "Data",
            "tables": [
                {
                    "data_range": "A1:C3",
                    "table_style": "TableStyleMedium9",
                }
            ],
        }
        with tempfile.TemporaryDirectory() as temporary:
            apply_workspace_recipe(
                temporary,
                self.make_recipe(),
                require_empty=True,
            )
            result = evaluate_workspace(temporary, _evaluator(assertion))
            wrong_style = dict(assertion)
            wrong_style["tables"] = [
                {
                    "data_range": "A1:C3",
                    "table_style": "TableStyleLight9",
                }
            ]
            wrong = evaluate_workspace(temporary, _evaluator(wrong_style))
        self.assertTrue(result["passed"], result)
        self.assertFalse(wrong["passed"], wrong)

    def test_chart_range_rejects_a_missing_middle_series(self):
        with tempfile.TemporaryDirectory() as temporary:
            apply_workspace_recipe(
                temporary,
                self.make_recipe(),
                require_empty=True,
            )
            from openpyxl import load_workbook

            path = Path(temporary) / "reports/sales.xlsx"
            workbook = load_workbook(path)
            chart = workbook["Data"]._charts[0]
            chart.ser = [chart.ser[-1]]
            workbook.save(path)
            workbook.close()
            result = evaluate_workspace(
                temporary,
                _evaluator(
                    {
                        "op": "xlsx_chart_properties",
                        "path": "reports/sales.xlsx",
                        "sheet": "Data",
                        "index": 0,
                        "properties": self.chart_properties(),
                    }
                ),
            )
        self.assertFalse(result["passed"], result)

    def test_chart_range_rejects_a_duplicate_series(self):
        with tempfile.TemporaryDirectory() as temporary:
            apply_workspace_recipe(
                temporary,
                self.make_recipe(),
                require_empty=True,
            )
            from copy import copy
            from openpyxl import load_workbook

            path = Path(temporary) / "reports/sales.xlsx"
            workbook = load_workbook(path)
            chart = workbook["Data"]._charts[0]
            chart.ser = [*chart.ser, copy(chart.ser[-1])]
            workbook.save(path)
            workbook.close()
            result = evaluate_workspace(
                temporary,
                _evaluator(
                    {
                        "op": "xlsx_chart_properties",
                        "path": "reports/sales.xlsx",
                        "sheet": "Data",
                        "index": 0,
                        "properties": self.chart_properties(),
                    }
                ),
            )
        self.assertFalse(result["passed"], result)

    def test_background_color_requires_a_solid_fill(self):
        with tempfile.TemporaryDirectory() as temporary:
            apply_workspace_recipe(
                temporary,
                self.make_recipe(),
                require_empty=True,
            )
            from openpyxl import load_workbook
            from openpyxl.styles import PatternFill

            path = Path(temporary) / "reports/sales.xlsx"
            workbook = load_workbook(path)
            workbook["Data"]["A1"].fill = PatternFill(
                fill_type=None,
                fgColor="4472C4",
            )
            workbook.save(path)
            workbook.close()
            result = evaluate_workspace(
                temporary,
                _evaluator(
                    {
                        "op": "xlsx_cell_style",
                        "path": "reports/sales.xlsx",
                        "sheet": "Data",
                        "cell": "A1",
                        "properties": {"bg_color": "4472C4"},
                    }
                ),
            )
        self.assertFalse(result["passed"], result)

    def test_strict_table_properties_reject_wrong_range_or_style(self):
        assertions = (
            {
                "op": "xlsx_table_properties",
                "path": "reports/sales.xlsx",
                "sheet": "Data",
                "table_name": "SalesTable",
                "data_range": "A1:B3",
                "table_style": "TableStyleMedium9",
            },
            {
                "op": "xlsx_table_properties",
                "path": "reports/sales.xlsx",
                "sheet": "Data",
                "table_name": "SalesTable",
                "data_range": "A1:C3",
                "table_style": "TableStyleMedium8",
            },
        )
        with tempfile.TemporaryDirectory() as temporary:
            apply_workspace_recipe(
                temporary,
                self.make_recipe(),
                require_empty=True,
            )
            for assertion in assertions:
                with self.subTest(assertion=assertion):
                    result = evaluate_workspace(
                        temporary,
                        _evaluator(assertion),
                    )
                    self.assertFalse(result["passed"], result)

    def test_each_strict_chart_property_is_enforced(self):
        wrong_values = {
            "chart_type": "line",
            "target_cell": "F2",
            "title": "Wrong title",
            "x_axis": "Wrong x axis",
            "y_axis": "Wrong y axis",
            "data_range": "A1:B3",
        }
        with tempfile.TemporaryDirectory() as temporary:
            apply_workspace_recipe(
                temporary,
                self.make_recipe(),
                require_empty=True,
            )
            for property_name, wrong_value in wrong_values.items():
                properties = self.chart_properties()
                properties[property_name] = wrong_value
                assertion = {
                    "op": "xlsx_chart_properties",
                    "path": "reports/sales.xlsx",
                    "sheet": "Data",
                    "index": 0,
                    "properties": properties,
                }
                with self.subTest(property_name=property_name):
                    result = evaluate_workspace(
                        temporary,
                        _evaluator(assertion),
                    )
                    self.assertFalse(result["passed"], result)

    def test_strict_chart_index_and_merged_range_set_are_enforced(self):
        assertions = (
            {
                "op": "xlsx_chart_properties",
                "path": "reports/sales.xlsx",
                "sheet": "Data",
                "index": 1,
                "properties": self.chart_properties(),
            },
            {
                "op": "xlsx_merged_ranges_equal",
                "path": "reports/sales.xlsx",
                "sheet": "Data",
                "ranges": [],
            },
            {
                "op": "xlsx_merged_ranges_equal",
                "path": "reports/sales.xlsx",
                "sheet": "Data",
                "ranges": ["A5:C5", "E5:F5"],
            },
        )
        with tempfile.TemporaryDirectory() as temporary:
            apply_workspace_recipe(
                temporary,
                self.make_recipe(),
                require_empty=True,
            )
            self._add_declared_merge(temporary)
            for assertion in assertions:
                with self.subTest(operation=assertion["op"]):
                    result = evaluate_workspace(
                        temporary,
                        _evaluator(assertion),
                    )
                    self.assertFalse(result["passed"], result)


@unittest.skipUnless(
    PDF_RUNTIME_AVAILABLE,
    "ReportLab and PyMuPDF are installed only in the Toolathlon runtime",
)
class PdfEvaluatorExtensionTests(unittest.TestCase):
    def test_page_count_and_one_indexed_page_text(self):
        recipe = {
            "directories": ["reports"],
            "files": [
                {
                    "path": "reports/brief.pdf",
                    "format": "pdf",
                    "content": [
                        "First page contains the overview.",
                        "Second page contains the approval code ZETA-42.",
                    ],
                }
            ],
            "remove": [],
        }
        evaluator = _evaluator(
            {
                "op": "pdf_page_count",
                "path": "reports/brief.pdf",
                "count": 2,
            },
            {
                "op": "pdf_page_text_contains",
                "path": "reports/brief.pdf",
                "page": 2,
                "value": "ZETA-42",
            },
        )
        with tempfile.TemporaryDirectory() as temporary:
            apply_workspace_recipe(temporary, recipe, require_empty=True)
            result = evaluate_workspace(temporary, evaluator)
            wrong_page = evaluate_workspace(
                temporary,
                _evaluator(
                    {
                        "op": "pdf_page_text_contains",
                        "path": "reports/brief.pdf",
                        "page": 1,
                        "value": "ZETA-42",
                    }
                ),
            )
        self.assertTrue(result["passed"], result)
        self.assertFalse(wrong_page["passed"])

    def test_exact_page_text_rejects_added_content(self):
        recipe = {
            "directories": ["reports"],
            "files": [
                {
                    "path": "reports/brief.pdf",
                    "format": "pdf",
                    "content": ["Approval code ZETA-42. Unexpected appendix."],
                }
            ],
            "remove": [],
        }
        with tempfile.TemporaryDirectory() as temporary:
            apply_workspace_recipe(temporary, recipe, require_empty=True)
            exact = evaluate_workspace(
                temporary,
                _evaluator(
                    {
                        "op": "pdf_page_text_equals",
                        "path": "reports/brief.pdf",
                        "page": 1,
                        "value": "Approval code ZETA-42. Unexpected appendix.",
                    }
                ),
            )
            missing_extra = evaluate_workspace(
                temporary,
                _evaluator(
                    {
                        "op": "pdf_page_text_equals",
                        "path": "reports/brief.pdf",
                        "page": 1,
                        "value": "Approval code ZETA-42.",
                    }
                ),
            )
        self.assertTrue(exact["passed"], exact)
        self.assertFalse(missing_extra["passed"], missing_extra)


class EvaluatorExtensionValidationTests(unittest.TestCase):
    def make_task(self, assertion):
        return {
            "initial_workspace": {},
            "oracle_final_state": {},
            "evaluator": _evaluator(assertion),
        }

    def test_pdf_page_is_one_indexed(self):
        for operation in ("pdf_page_text_contains", "pdf_page_text_equals"):
            with self.subTest(operation=operation):
                with self.assertRaisesRegex(ValueError, "positive 1-indexed page"):
                    validate_task_spec_structure(
                        self.make_task(
                            {
                                "op": operation,
                                "path": "brief.pdf",
                                "page": 0,
                                "value": "text",
                            }
                        )
                    )

    def test_new_assertions_retain_workspace_path_safety(self):
        for assertion in (
            {
                "op": "xlsx_sheet_names_equal",
                "path": "../outside.xlsx",
                "sheets": ["Data"],
            },
            {
                "op": "pdf_page_count",
                "path": "/outside.pdf",
                "count": 1,
            },
        ):
            with self.subTest(operation=assertion["op"]):
                with self.assertRaises(ValueError):
                    validate_task_spec_structure(self.make_task(assertion))

    def test_style_properties_are_restricted_to_declarative_fields(self):
        with self.assertRaisesRegex(ValueError, "unsupported properties"):
            validate_task_spec_structure(
                self.make_task(
                    {
                        "op": "xlsx_cell_style",
                        "path": "report.xlsx",
                        "sheet": "Data",
                        "cell": "A1",
                        "properties": {"callback": "run_code()"},
                    }
                )
            )

    def test_strict_table_schema_rejects_extra_fields(self):
        with self.assertRaisesRegex(ValueError, "requires exactly"):
            validate_task_spec_structure(
                self.make_task(
                    {
                        "op": "xlsx_table_properties",
                        "path": "report.xlsx",
                        "sheet": "Data",
                        "table_name": "SalesTable",
                        "data_range": "A1:C3",
                        "table_style": "TableStyleMedium9",
                        "ignored": "would weaken the assertion",
                    }
                )
            )

    def test_strict_chart_schema_requires_all_six_properties(self):
        with self.assertRaisesRegex(ValueError, "requires exactly chart_type"):
            validate_task_spec_structure(
                self.make_task(
                    {
                        "op": "xlsx_chart_properties",
                        "path": "report.xlsx",
                        "sheet": "Data",
                        "index": 0,
                        "properties": {
                            "chart_type": "bar",
                            "target_cell": "E2",
                            "title": "Sales",
                            "x_axis": "Region",
                            "data_range": "A1:C3",
                        },
                    }
                )
            )

    def test_chart_index_and_merged_ranges_are_strictly_validated(self):
        invalid_assertions = (
            {
                "op": "xlsx_chart_properties",
                "path": "report.xlsx",
                "sheet": "Data",
                "index": -1,
                "properties": {
                    "chart_type": "bar",
                    "target_cell": "E2",
                    "title": "Sales",
                    "x_axis": "Region",
                    "y_axis": "Value",
                    "data_range": "A1:C3",
                },
            },
            {
                "op": "xlsx_merged_ranges_equal",
                "path": "report.xlsx",
                "sheet": "Data",
                "ranges": ["A1:C1", "a1:c1"],
            },
        )
        for assertion in invalid_assertions:
            with self.subTest(operation=assertion["op"]):
                with self.assertRaises(ValueError):
                    validate_task_spec_structure(self.make_task(assertion))


if __name__ == "__main__":
    unittest.main()
